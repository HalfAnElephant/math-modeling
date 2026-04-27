from pathlib import Path

import pytest

from c_uav_inspection.data import _read_matrix_sheet, _read_targets, _read_uav_params, load_problem_data, validate_problem_data


DATA_PATH = Path("2026同济数学建模竞赛赛题/2026C数据.xlsx")


def test_load_problem_data_counts_targets_and_params():
    data = load_problem_data(DATA_PATH)

    assert len(data.targets) == 16
    assert data.params.k_max == 4
    assert data.params.effective_energy_limit_j == 135000
    assert data.params.hover_power_j_per_s == 220
    assert data.params.battery_swap_time_s == 300
    assert data.params.operating_horizon_s == 2600


def test_loaded_targets_preserve_key_values():
    data = load_problem_data(DATA_PATH)
    by_id = {target.node_id: target for target in data.targets}

    assert by_id[1].node_name == "B1-东立面裂缝"
    assert by_id[1].base_hover_time_s == 50
    assert by_id[1].direct_confirm_time_s == 200
    assert by_id[16].priority_weight == 3
    assert by_id[16].direct_confirm_time_s == 560


def test_matrices_include_depot_and_target_pairs():
    data = load_problem_data(DATA_PATH)

    assert data.flight_time_s[(0, 1)] > 0
    assert data.flight_energy_j[(0, 16)] > 0
    assert data.ground_time_s[("P0", "MP01")] > 0
    assert data.flight_time_s[(3, 3)] == 0


def test_read_uav_params_raises_on_unexpected_parameter():
    """CR-004 regression: _read_uav_params must reject extra parameters."""
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active

    # 14 params (rows 4-17), with one expected key replaced by an unexpected one
    params = [
        ("K_max", 4),
        ("battery_capacity_J", 500000),
        ("safety_reserve_J", 80000),
        ("effective_energy_limit_J", 135000),
        ("horizontal_speed_mps", 20),
        ("vertical_speed_mps", 4),
        ("horizontal_energy_J_per_m", 12),
        ("up_energy_J_per_m", 15),
        ("down_energy_J_per_m", 8),
        ("hover_power_J_per_s", 220),
        ("battery_swap_time_s", 300),
        ("operating_horizon_s", 2600),
        ("walking_speed_mps", 1.2),
        ("UNEXPECTED_EXTRA_PARAM", 42),  # replaces walking_detour_factor
    ]

    for i, (name, val) in enumerate(params, start=4):
        ws.cell(row=i, column=1, value=name)
        ws.cell(row=i, column=2, value=val)

    with pytest.raises(ValueError, match=r"(?i)(unexpected|unknown|extra)"):
        _read_uav_params(ws)


def test_read_uav_params_raises_on_missing_parameter():
    """CR-004 regression: _read_uav_params must reject missing parameters."""
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active

    params = [
        ("K_max", 4),
        ("battery_capacity_J", 500000),
        ("safety_reserve_J", 80000),
        ("effective_energy_limit_J", 135000),
        ("horizontal_speed_mps", 20),
        ("vertical_speed_mps", 4),
        ("horizontal_energy_J_per_m", 12),
        ("up_energy_J_per_m", 15),
        ("down_energy_J_per_m", 8),
        ("hover_power_J_per_s", 220),
        ("battery_swap_time_s", 300),
        ("operating_horizon_s", 2600),
        # walking_speed_mps and walking_detour_factor intentionally missing
    ]

    for i, (name, val) in enumerate(params, start=4):
        ws.cell(row=i, column=1, value=name)
        ws.cell(row=i, column=2, value=val)

    with pytest.raises(ValueError, match=r"(?i)(missing)"):
        _read_uav_params(ws)


def test_validate_problem_data_returns_expected_summary():
    data = load_problem_data(DATA_PATH)
    summary = validate_problem_data(data)

    assert summary["target_count"] == 16
    assert summary["base_hover_sum_s"] == 790
    assert summary["direct_hover_sum_s"] == 5210
    assert summary["confirm_thresholds_valid"] is True
    assert summary["max_single_direct_confirm_energy_j"] <= data.params.effective_energy_limit_j


def _make_target_row(
    node_id: int = 1,
    node_name: str = "B1-东立面裂缝",
    building_id: str = "B1",
    x_m: float = 10.0,
    y_m: float = 20.0,
    z_m: float = 5.0,
    priority_level: str = "高",
    priority_weight: int = 5,
    issue_type: str = "裂缝",
    base_hover_time_s: float = 50.0,
    direct_confirm_time_s: float = 200.0,
    manual_point_id: str = "MP01",
    manual_x_m: float = 15.0,
    manual_y_m: float = 25.0,
    manual_service_time_s: float = 60.0,
) -> list:
    """Build a single row matching the NodeData column layout.

    Column index 11 (extra_confirm_time_s) is intentionally skipped in _read_targets,
    so it is placed as None here to match the actual Excel structure.
    """
    return [
        node_id,
        node_name,
        building_id,
        x_m,
        y_m,
        z_m,
        priority_level,
        priority_weight,
        issue_type,
        base_hover_time_s,
        direct_confirm_time_s,
        None,  # column 11: extra_confirm_time_s (skipped)
        manual_point_id,
        manual_x_m,
        manual_y_m,
        manual_service_time_s,
    ]


def test_read_targets_skips_none_rows():
    """CR-005: _read_targets must skip None rows without TypeError.

    When empty rows exist within the data range, _read_targets should not
    crash with TypeError (as the old code did). Instead, it should either
    skip them gracefully and proceed, or detect the resultant wrong count
    and raise a clean ValueError. Both outcomes are better than a cryptic
    TypeError from int(None).
    """
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active

    # 14 valid rows + 2 empty rows within the data range.  After skipping
    # the empty rows, the count validation should catch the discrepancy.
    rows_data = []
    for i in range(1, 17):
        if i in (3, 8):
            rows_data.append(None)
        else:
            rows_data.append(_make_target_row(node_id=i))

    for idx, row_data in enumerate(rows_data, start=5):
        if row_data is None:
            continue  # leave the row empty
        for col_idx, value in enumerate(row_data, start=1):
            ws.cell(row=idx, column=col_idx, value=value)

    with pytest.raises(ValueError, match=r"(?i)(expected 16|target count|16 target)"):
        _read_targets(ws)


def test_read_targets_raises_on_too_few_targets():
    """CR-005: _read_targets must raise ValueError when target count != 16."""
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active

    # Only 14 target rows when we expect 16
    for i in range(1, 15):
        row_data = _make_target_row(node_id=i)
        for col_idx, value in enumerate(row_data, start=1):
            ws.cell(row=i + 4, column=col_idx, value=value)

    with pytest.raises(ValueError, match=r"(?i)(expected 16|target count|16 target)"):
        _read_targets(ws)


def test_read_targets_raises_on_too_many_targets():
    """CR-005: _read_targets must raise ValueError when target count > 16."""
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active

    # 18 target rows
    for i in range(1, 19):
        row_data = _make_target_row(node_id=i)
        for col_idx, value in enumerate(row_data, start=1):
            ws.cell(row=i + 4, column=col_idx, value=value)

    with pytest.raises(ValueError, match=r"(?i)(expected 16|target count|16 target)"):
        _read_targets(ws)


def test_read_targets_with_empty_rows_still_validates_count():
    """CR-005: After skipping None rows, count must still be validated.

    If empty rows mask the fact that only 14 actual target rows exist,
    _read_targets must still detect the wrong count.
    """
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active

    # 14 real rows + 2 empty rows in range 5-20 = only 14 real targets
    for i in range(1, 15):
        row_data = _make_target_row(node_id=i)
        for col_idx, value in enumerate(row_data, start=1):
            ws.cell(row=i + 4, column=col_idx, value=value)
    # Leave rows 19 and 20 empty

    with pytest.raises(ValueError, match=r"(?i)(expected 16|target count|16 target)"):
        _read_targets(ws)


# ── CR-006: _read_matrix_sheet empty-row validation ───────────────────────


def _make_matrix_header_row(col_ids: list[int]) -> list:
    """Build a header row (row 3) matching matrix sheet column layout.

    Column index 0 = "From \\ Node", column 1 = label (e.g. "节点0"),
    columns 2+ = node IDs.
    """
    return [None, None] + col_ids


def _make_matrix_data_row(from_id: int, num_cols: int) -> list:
    """Build a data row: [from_id, label, v0, v1, ..., vN]."""
    return [from_id, f"节点{from_id}"] + [float(from_id * 100 + c) for c in range(num_cols)]


def test_read_matrix_sheet_skips_none_rows():
    """CR-006: _read_matrix_sheet must skip None rows without TypeError.

    When a row has None in the row identifier column but non-None data
    in other columns (e.g. a partially malformed row), _read_matrix_sheet
    should skip the row gracefully instead of crashing with TypeError.
    """
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active

    col_ids = [0, 1, 2]
    header = _make_matrix_header_row(col_ids)
    for col_idx, value in enumerate(header, start=1):
        ws.cell(row=3, column=col_idx, value=value)

    # Write rows: from_id=0, a row with None in col 1 (row_id), then
    # from_id=1 and from_id=2.  The None-row must have data in other
    # columns so it is NOT detected as a fully-empty "end of data" row.
    rows: list[tuple[int, list]] = [
        (4, _make_matrix_data_row(0, len(col_ids))),
        # row 5: None row_id but data in other columns
        (5, [None, "label"] + [999.0, 999.0, 999.0]),
        (6, _make_matrix_data_row(1, len(col_ids))),
        (7, _make_matrix_data_row(2, len(col_ids))),
    ]
    for row_num, row_data in rows:
        for col_idx, value in enumerate(row_data, start=1):
            ws.cell(row=row_num, column=col_idx, value=value)

    result = _read_matrix_sheet(ws, key_type="int")

    # Must contain data from the valid rows
    assert result[(0, 0)] == 0.0
    assert result[(1, 0)] == 100.0
    assert result[(2, 0)] == 200.0
    # Must NOT contain data from the None-row
    assert not any(key[0] is None for key in result)


def test_read_matrix_sheet_breaks_on_fully_empty_row():
    """CR-006: _read_matrix_sheet must stop when encountering a fully empty row.

    A fully empty row (all cells None) signals end-of-data, just as
    _read_targets does with its all-None check.
    """
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active

    col_ids = [0, 1, 2]
    header = _make_matrix_header_row(col_ids)
    for col_idx, value in enumerate(header, start=1):
        ws.cell(row=3, column=col_idx, value=value)

    # Write one data row, then leave all subsequent rows fully empty.
    row_data = _make_matrix_data_row(0, len(col_ids))
    for col_idx, value in enumerate(row_data, start=1):
        ws.cell(row=4, column=col_idx, value=value)

    # Rows 5-20 are all empty (not written)

    result = _read_matrix_sheet(ws, key_type="int")

    # Should only contain node 0 entries
    assert len(result) == len(col_ids)
    assert all(key[0] == 0 for key in result)
    assert (0, 0) in result
